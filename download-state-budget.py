#!/usr/bin/env python3
"""Download the settled State budget and write data/spain-budget.csv.

The Comptroller's liquidation of the State budget is the annual budget law
after it has been spent. Health, education, care, housing, and transport are
the same policies as in a community liquidation. Interest is the financial
chapter. Pensions are the State budget's pensions policy, mostly civil-service
pensions. Repayment of principal and transfers inside the State budget stay
out of the total. Every other spending policy is kept on its own. Roads and
rail inside a mixed infrastructure policy are counted with transport, and what
remains of that policy stays on its own bar. The series starts in 2002.
"""

from __future__ import annotations

import csv
import os
import re
import subprocess
import sys
import tempfile
import unicodedata
import urllib.parse
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "data" / "spain-budget.csv"
PAGE = (
    "https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/"
    "ContabilidadPublica/CPE/EjecucionPresupuestaria/Paginas/ialiquidacionestado.aspx"
)
ORIGIN = "https://www.igae.pap.hacienda.gob.es"
USER_AGENT = "plan-mery/state-budget"
NAMED = ("health", "education", "care", "housing", "transport", "interest", "pensions")
TARGETS = (
    "justice", "defence", "security", "foreign", "employment", "labour",
    "social_security", "culture", "agriculture", "farm_works", "industry", "energy",
    "commerce", "tourism", "research", "economic", "leadership", "general",
    "statistics", "tax", "transfers", "eu_transfers", "communications",
    "community", "other_social", "resilient",
)
COLUMNS = (*NAMED, *TARGETS, "rest", "spent")
PARTS_PATH = ROOT / "data" / "budget-parts.csv"
# Program groups inside four policies. Codes changed in 2005: before that, 111 was
# the head of state and 142 was the courts; afterwards 911 is the head of state and 142 is embassies.
NEW_GROUP = {
    "911": "leadership",
    "912": "leadership",
    "921": "general_admin",
    "922": "general_territory",
    "923": "general_finance",
    "924": "general_elections",
    "929": "general_contingency",
    "111": "justice_admin",
    "112": "courts",
    "113": "justice_other",
    "141": "foreign_other",
    "142": "embassies",
    "143": "cooperation",
    "144": "cooperation",
}
OLD_GROUP = {
    "111": "leadership",
    "112": "leadership",
    "121": "general_admin",
    "124": "general_territory",
    "126": "general_other",
    "131": "foreign_other",
    "132": "embassies",
    "134": "cooperation",
    "141": "justice_admin",
    "142": "courts",
    "146": "justice_other",
}
PART_PARENT = {
    "crown": "leadership",
    "parliament": "leadership",
    "audit": "leadership",
    "constitutional": "leadership",
    "council": "leadership",
    "government": "leadership",
    "leadership_other": "leadership",
    "courts": "justice",
    "judicial_council": "justice",
    "justice_other": "justice",
    "embassies": "foreign",
    "cooperation": "foreign",
    "foreign_other": "foreign",
    "general_admin": "general",
    "general_territory": "general",
    "general_finance": "general",
    "general_mutual": "general",
    "general_purchase": "general",
    "general_elections": "general",
    "general_contingency": "general",
    "general_memory": "general",
    "general_other": "general",
    "staff": "general",
    "goods": "general",
    "current_transfers": "general",
    "capital_transfers": "general",
    "investment": "general",
    "assets": "general",
    "contingency": "general",
}
NUMBER = re.compile(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?|\d+(?:,\d+)?")
POLICY_LINE = re.compile(
    r"^\s*(\d+)\.\s+(.+?)\s{2,}(.*)$"
)
YEAR = re.compile(r"(20\d{2}|19\d{2})")


def ensure_readers():
    venv = Path(tempfile.gettempdir()) / "plan-mery-xlsvenv"
    python = venv / "bin" / "python"
    if not python.exists():
        subprocess.check_call([sys.executable, "-m", "venv", str(venv)])
    if subprocess.run([str(python), "-c", "import openpyxl, xlrd"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL).returncode != 0:
        subprocess.check_call([str(python), "-m", "pip", "install", "openpyxl", "xlrd==1.2.0"])
    if Path(sys.prefix).resolve() != venv.resolve():
        os.execv(str(python), [str(python), __file__, *sys.argv[1:]])


def fold(value: str) -> str:
    text = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return " ".join(text.lower().split())


def spanish_number(token: str) -> float:
    return float(token.replace(".", "").replace(",", "."))


def bucket(name: str) -> str:
    folded = fold(name)
    if "transferencias internas" in folded:
        return "internal"
    if "deuda publica" in folded:
        return "debt"
    if re.search(r"\bpensiones\b", folded) or "seguridad y proteccion social" in folded:
        return "pensions"
    if re.search(r"\bsanidad\b", folded):
        return "health"
    if re.search(r"\beducacion\b", folded):
        return "education"
    if "servicios sociales" in folded or folded.strip() == "promocion social" or folded.endswith(" promocion social"):
        return "care"
    if "vivienda" in folded:
        return "housing"
    if "agraria" in folded and "infraestructura" in folded:
        return "farm_works"
    if "transporte" in folded:
        return "transport"
    if "infraestructura" in folded and "ecosistema" not in folded and "resiliente" not in folded:
        return "transport"
    if "ecosistema" in folded or "resiliente" in folded:
        return "resilient"
    rules = (
        ("justicia", "justice"),
        ("defensa", "defence"),
        ("seguridad ciudadana", "security"),
        ("proteccion civil", "security"),
        ("politica exterior", "foreign"),
        ("relaciones exteriores", "foreign"),
        ("fomento del empleo", "employment"),
        ("trabajo y economia", "labour"),
        ("seguridad social", "social_security"),
        ("inclusion", "social_security"),
        ("agricultura", "agriculture"),
        ("cultura", "culture"),
        ("industria", "industry"),
        ("energia", "energy"),
        ("mineria", "energy"),
        ("comercio", "commerce"),
        ("turismo", "tourism"),
        ("investigac", "research"),
        ("caracter economico", "economic"),
        ("regulacion economica", "economic"),
        ("alta direccion", "leadership"),
        ("organos constitucionales", "leadership"),
        ("servicios de caracter general", "general"),
        ("administracion general", "general"),
        ("estadistica", "statistics"),
        ("administracion financiera", "tax"),
        ("regulacion financiera", "tax"),
        ("union europea", "eu_transfers"),
        ("transferencias a", "transfers"),
        ("comunicaciones", "communications"),
        ("bienestar comunitario", "community"),
        ("otros servicios comunitarios", "other_social"),
    )
    for needle, key in rules:
        if needle in folded:
            return key
    return "rest"


def is_area(name: str) -> bool:
    letters = [char for char in name if char.isalpha()]
    return bool(letters) and not any(char.islower() for char in letters)


def numbers_of(text: str) -> list[float]:
    return [spanish_number(token) for token in NUMBER.findall(text)]


def compose(policies: dict[str, float], interest: float, principal: float, roads: float) -> dict[str, float]:
    gross = sum(policies.values())
    base = gross - policies.get("internal", 0) - principal
    named = {
        "health": policies.get("health", 0),
        "education": policies.get("education", 0),
        "care": policies.get("care", 0),
        "housing": policies.get("housing", 0),
        "transport": policies.get("transport", 0) + roads,
        "interest": interest,
        "pensions": policies.get("pensions", 0),
    }
    targets = {key: policies.get(key, 0.0) for key in TARGETS}
    targets["resilient"] -= roads
    rest = base - sum(named.values()) - sum(targets.values())
    if "internal" not in policies or "health" not in policies or "pensions" not in policies or "debt" not in policies:
        raise RuntimeError(f"Policy table is missing a required policy: {sorted(policies)}")
    # The debt policy and the interest-plus-principal chapters can differ by a small amount.
    if targets["resilient"] < -1 or rest < -0.01 * base or min(base, interest, named["health"], named["pensions"]) <= 0:
        raise RuntimeError(f"Budget parts are not usable: base {base}, rest {rest}, resilient {targets['resilient']}")
    if rest < 0:
        rest = 0
    spent = dict(named)
    spent.update(targets)
    spent["rest"] = max(rest, 0)
    spent["spent"] = base
    return {name: round(spent[name], 2) for name in COLUMNS}


def sheet_title(rows: list[tuple]) -> str:
    for row in rows[:4]:
        for cell in row:
            if cell and "LIQUID" in str(cell).upper():
                return fold(str(cell))
    return ""


def sheet_rows(book, name: str) -> list[tuple]:
    sheet = book[name]
    return [tuple(row) for row in sheet.iter_rows(values_only=True)]


def open_book(path: Path):
    if path.suffix.lower() == ".xls":
        import xlrd
        book = xlrd.open_workbook(path)
        class Adapter:
            sheetnames = book.sheet_names()
            def __getitem__(self, name):
                return book.sheet_by_name(name)
        class Sheet:
            def __init__(self, raw):
                self.raw = raw
            def iter_rows(self, values_only=True):
                for index in range(self.raw.nrows):
                    yield tuple(self.raw.cell_value(index, col) for col in range(self.raw.ncols))
        class Books:
            sheetnames = book.sheet_names()
            def __getitem__(self, name):
                return Sheet(book.sheet_by_name(name))
        return Books()
    import openpyxl
    return openpyxl.load_workbook(path, data_only=True, read_only=True)


def find_sheet(book, phrase: str) -> list[tuple]:
    found = None
    for name in book.sheetnames:
        rows = sheet_rows(book, name)
        if phrase in sheet_title(rows):
            found = rows
    if found is None:
        raise RuntimeError(f"Workbook has no sheet {phrase}")
    return found


def obligation_column(rows: list[tuple]) -> int:
    for row in rows[:6]:
        for index, cell in enumerate(row):
            if cell and "obligaciones reconocidas" in fold(str(cell)):
                return index
    raise RuntimeError("Workbook has no recognized-obligations column")


def amount(cell) -> float | None:
    if cell is None or cell == "":
        return None
    if isinstance(cell, (int, float)):
        return float(cell)
    text = str(cell).strip()
    if not text or text == "-":
        return None
    return spanish_number(text)


def policies_from_sheet(rows: list[tuple]) -> dict[str, float]:
    column = obligation_column(rows)
    policies: dict[str, float] = {}
    for row in rows:
        if not row or row[0] is None:
            continue
        label = str(row[0])
        if not label[:1].isspace():
            continue
        name = label.strip()
        if is_area(name):
            continue
        value = amount(row[column] if column < len(row) else None)
        if value is None:
            continue
        kind = bucket(name.split(".", 1)[-1])
        policies[kind] = policies.get(kind, 0) + value
    if "health" not in policies or "internal" not in policies:
        raise RuntimeError("Policy sheet is missing health or internal transfers")
    return policies


def chapter_amount(rows: list[tuple], marker: str) -> float:
    column = obligation_column(rows)
    found = None
    for row in rows:
        if not row or row[0] is None:
            continue
        label = fold(str(row[0]))
        if label.startswith(marker):
            found = amount(row[column] if column < len(row) else None)
    if found is None:
        raise RuntimeError(f"Chapter sheet is missing {marker}")
    return found


def transport_programs(rows: list[tuple]) -> float:
    column = obligation_column(rows)
    total = 0.0
    for row in rows:
        if not row or row[0] is None:
            continue
        label = str(row[0]).strip()
        if not label or is_area(label):
            continue
        folded = fold(label)
        if not any(word in folded for word in ("carretera", "ferrocarril", "ferroviario", "red nacional de transporte", "red transeuropea")):
            continue
        value = amount(row[column] if column < len(row) else None)
        if value:
            total += value
    return total


def parse_workbook(path: Path) -> dict[str, float]:
    book = open_book(path)
    policies = policies_from_sheet(find_sheet(book, "situacion por areas de gasto y politicas"))
    chapters = find_sheet(book, "resumen por capitulos")
    interest = chapter_amount(chapters, "3.")
    principal = chapter_amount(chapters, "9.")
    roads = transport_programs(find_sheet(book, "situacion de los programas")) if policies.get("resilient") else 0.0
    return compose(policies, interest, principal, roads)


def parse_pdf(path: Path) -> dict[str, float]:
    text = subprocess.check_output(["pdftotext", "-layout", str(path), "-"], text=True, errors="replace")
    chunk = functional_chunk(text, path.name)
    policies: dict[str, float] = {}
    for line in chunk.splitlines():
        match = POLICY_LINE.match(line)
        if not match:
            continue
        values = numbers_of(match.group(3))
        if len(values) != 6:
            continue
        name = match.group(2).strip()
        if is_area(name):
            continue
        kind = bucket(name)
        policies[kind] = policies.get(kind, 0) + values[1]
    interest = last_chapter(text, "3. GASTOS FINANCIEROS")
    principal = last_chapter(text, "9. PASIVOS FINANCIEROS")
    if "health" not in policies:
        raise RuntimeError(f"{path.name} policy table is missing health")
    return compose(policies, interest, principal, 0.0)


def functional_chunk(text: str, name: str) -> str:
    markers = (
        "SITUACIÓN POR ÁREAS DE GASTO Y POLÍTICAS",
        "SITUACIÓN POR GRUPOS Y FUNCIONES",
    )
    start = 0
    upper = text.upper()
    while start < len(text):
        indexes = [upper.find(marker, start) for marker in markers]
        indexes = [index for index in indexes if index >= 0]
        if not indexes:
            break
        index = min(indexes)
        end = re.search(r"\n\s*TOTALES\b", text[index:])
        if end:
            chunk = text[index:index + end.start()]
            if re.search(r"1\.\s+Sanidad", chunk):
                return chunk
        start = index + 20
    raise RuntimeError(f"{name} has no functional heading")


def last_chapter(text: str, heading: str) -> float:
    found = None
    for line in text.splitlines():
        if heading not in line.upper():
            continue
        tail = line.upper().split(heading, 1)[-1]
        values = numbers_of(tail)
        if len(values) == 6:
            found = values[1]
    if found is None:
        raise RuntimeError(f"Missing chapter {heading}")
    return found


def links() -> dict[int, str]:
    request = urllib.request.Request(PAGE, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=60) as response:
        html = response.read().decode("utf-8", "replace")
    chosen: dict[int, tuple[int, str]] = {}
    for href in re.findall(r'href="([^"]+)"', html, re.I):
        name = urllib.parse.unquote(href).lower()
        if "anexo" in name or "ooaa" in name or "explot" in name:
            continue
        if "liquida" not in name:
            continue
        if not name.endswith((".pdf", ".xls", ".xlsx")):
            continue
        match = YEAR.search(urllib.parse.unquote(href))
        if not match:
            continue
        year = int(match.group(1))
        rank = 2 if name.endswith(".xlsx") else 1 if name.endswith(".xls") else 0
        url = href if href.startswith("http") else ORIGIN + href
        parts = urllib.parse.urlsplit(url)
        url = urllib.parse.urlunsplit((
            parts.scheme,
            parts.netloc,
            urllib.parse.quote(urllib.parse.unquote(parts.path), safe="/"),
            parts.query,
            "",
        ))
        if year not in chosen or rank > chosen[year][0]:
            chosen[year] = (rank, url)
    if 2002 not in chosen or max(chosen) < 2024 or any(year not in chosen for year in range(2002, max(chosen) + 1)):
        raise RuntimeError(f"Liquidation index is missing years: {sorted(chosen)}")
    return {year: url for year, (_rank, url) in chosen.items()}


def download(url: str, year: int) -> Path:
    suffix = Path(urllib.parse.urlparse(url).path).suffix or ".bin"
    target = Path(tempfile.gettempdir()) / f"plan-mery-state-{year}{suffix}"
    if not target.exists() or target.stat().st_size < 1000:
        request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
        with urllib.request.urlopen(request, timeout=120) as response:
            target.write_bytes(response.read())
    return target


def classify_group(code: str, name: str, old: bool) -> str | None:
    folded = fold(name)
    if code and code in (OLD_GROUP if old else NEW_GROUP):
        return (OLD_GROUP if old else NEW_GROUP)[code]
    if folded.startswith("alta direccion del estado") or folded.startswith("alta direccion del gobierno"):
        return "leadership"
    if folded.startswith("administracion de justicia"):
        return "courts"
    if folded.startswith("administracion general de justicia"):
        return "justice_admin"
    if folded == "registros" or folded.startswith("registros "):
        return "justice_other"
    if "accion diplomatica" in folded:
        return "embassies"
    if folded.startswith("cooperac"):
        return "cooperation"
    if "relaciones exteriores" in folded:
        return "foreign_other"
    if "polit" in folded and "ext" in folded:
        return "foreign_other"
    if not old and folded.startswith("imprev"):
        return "general_contingency"
    if not old and "participacion ciudadana" in folded:
        return "general_elections"
    if "memoria democratica" in folded:
        return "general_memory"
    if folded in {"administracion general", "servicios generales y funcion publica"}:
        return "general_admin"
    if folded.startswith("otros servicios generales"):
        return "general_other"
    if "organizacion territorial" in folded or "relativos a las aa" in folded:
        return "general_territory"
    if not old and "economia y hacienda" in folded:
        return "general_finance"
    if "caracter general" in folded:
        return "general_other"
    if "justicia" in folded:
        return "justice_other"
    return None


def refine_child(group_part: str, name: str) -> str:
    folded = fold(name)
    if group_part == "leadership":
        if "majestad" in folded or "jefatura del estado" in folded:
            return "crown"
        if "cortes generales" in folded:
            return "parliament"
        if "cuentas" in folded:
            return "audit"
        if "constitucional" in folded:
            return "constitutional"
        if "consejo de estado" in folded:
            return "council"
        if "presidencia" in folded:
            return "government"
        return "leadership_other"
    if group_part == "justice_admin":
        if "poder judicial" in folded:
            return "judicial_council"
        return "justice_other"
    if group_part == "general_finance":
        if "mutual" in folded:
            return "general_mutual"
        if "contratacion central" in folded:
            return "general_purchase"
        return "general_finance"
    return group_part


def parts_from_items(items: list[tuple[int, str, str, float]]) -> dict[str, float]:
    """items are (level, code, name, obligations). Level 0 is a program group."""
    groups = []
    current = None
    for level, code, name, value in items:
        if level == 0:
            current = {"code": code, "name": name, "amount": value, "children": []}
            groups.append(current)
        elif current is not None:
            current["children"].append((name, value))
    old = any(group["code"] == "111" and "alta direccion" in fold(group["name"]) for group in groups)
    parts: dict[str, float] = {}
    split = {"leadership", "justice_admin", "general_finance"}
    for group in groups:
        kind = classify_group(group["code"], group["name"], old)
        if kind is None:
            continue
        if kind in split and group["children"]:
            used = 0.0
            for name, value in group["children"]:
                key = refine_child(kind, name)
                parts[key] = parts.get(key, 0.0) + value
                used += value
            gap = group["amount"] - used
            if abs(gap) > 1:
                fallback = {"leadership": "leadership_other", "justice_admin": "justice_other", "general_finance": "general_finance"}[kind]
                parts[fallback] = parts.get(fallback, 0.0) + gap
        else:
            key = {"justice_admin": "justice_other", "leadership": "leadership_other"}.get(kind, kind)
            parts[key] = parts.get(key, 0.0) + group["amount"]
    return {key: round(value, 2) for key, value in parts.items() if value > 1}


def sheet_group_items(rows: list[tuple]) -> list[tuple[int, str, str, float]]:
    column = obligation_column(rows)
    items = []
    for row in rows:
        if not row or row[0] is None:
            continue
        raw = str(row[0])
        label = raw.strip()
        if not label:
            continue
        folded = fold(label)
        if folded.startswith("total") or "liquidacion" in folded or folded.startswith("miles"):
            continue
        value = amount(row[column] if column < len(row) else None)
        if value is None:
            continue
        match = re.match(r"^(\d{3})-(.+)$", label)
        code = match.group(1) if match else ""
        name = match.group(2).strip() if match else label
        level = 1 if raw[:1].isspace() else 0
        items.append((level, code, name, value))
    return items


def pdf_group_items(path: Path) -> list[tuple[int, str, str, float]]:
    text = subprocess.check_output(["pdftotext", "-layout", str(path), "-"], text=True, errors="replace")
    group_line = re.compile(r"^\s*(\d{3})-(.+?)\s{2,}(\d.*\d)\s*$")
    child_line = re.compile(r"^\s+(\D.*?)\s{2,}(\d.*\d)\s*$")
    items = []
    seen_group = False
    for line in text.splitlines():
        grouped = group_line.match(line)
        if grouped:
            values = [spanish_number(token) for token in NUMBER.findall(grouped.group(3))]
            if len(values) < 6:
                continue
            seen_group = True
            items.append((0, grouped.group(1), grouped.group(2).strip(), values[1]))
            continue
        if not seen_group:
            continue
        if fold(line).startswith("totales"):
            break
        child = child_line.match(line)
        if not child:
            continue
        values = [spanish_number(token) for token in NUMBER.findall(child.group(2))]
        if len(values) < 6:
            continue
        items.append((1, "", child.group(1).strip(), values[1]))
    if not items:
        raise RuntimeError(f"{path.name} has no program groups")
    return items


def state_parts(path: Path) -> dict[str, float]:
    if path.suffix.lower() == ".pdf":
        return parts_from_items(pdf_group_items(path))
    book = open_book(path)
    return parts_from_items(sheet_group_items(find_sheet(book, "situacion de los grupos de programas")))


def save_parts(replacement: dict[tuple[str, int], dict[str, float]]) -> None:
    """Replace the given slug-year rows in the parts file and keep every other row."""
    existing: dict[tuple[str, int, str], float] = {}
    if PARTS_PATH.exists():
        with PARTS_PATH.open(encoding="utf-8", newline="") as handle:
            for row in csv.DictReader(handle, delimiter=";"):
                key = (row["slug"], int(row["year"]), row["part"])
                existing[key] = float(row["amount"])
    for (slug, year), parts in replacement.items():
        for part in [key for key in list(existing) if key[0] == slug and key[1] == year]:
            existing.pop(part, None)
        for part, amount in parts.items():
            if part not in PART_PARENT:
                raise RuntimeError(f"Unknown budget part {part}")
            existing[(slug, year, part)] = amount
    PARTS_PATH.parent.mkdir(exist_ok=True)
    with PARTS_PATH.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(("slug", "year", "parent", "part", "amount"))
        for slug, year, part in sorted(existing):
            writer.writerow((slug, year, PART_PARENT[part], part, f"{existing[(slug, year, part)]:.2f}"))


def main() -> None:
    ensure_readers()
    OUT.parent.mkdir(exist_ok=True)
    rows = []
    part_rows = {}
    for year, url in sorted(links().items()):
        path = download(url, year)
        parsed = parse_workbook(path) if path.suffix.lower() in {".xls", ".xlsx"} else parse_pdf(path)
        rows.append((year, parsed))
        part_rows[("espana", year)] = state_parts(path)
        print(f"{year} health {parsed['health']:.0f} transport {parsed['transport']:.0f} interest {parsed['interest']:.0f} rest {parsed['rest']:.0f} spent {parsed['spent']:.0f}", file=sys.stderr)
        if parsed["rest"] > 1000:
            raise RuntimeError(f"{year} still has an unclassified remainder of {parsed['rest']:.0f}")
    with OUT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(("slug", "year", *COLUMNS))
        for year, parsed in rows:
            writer.writerow(("espana", year, *(f"{parsed[name]:.2f}" for name in COLUMNS)))
    save_parts(part_rows)
    print(f"Wrote {OUT} ({len(rows)} years)")


if __name__ == "__main__":
    main()
