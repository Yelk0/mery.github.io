#!/usr/bin/env python3
"""Download Spain's public-accounts balance and write data/spain.csv.

IGAE publishes the non-financial accounts of all public administrations
together. B.9 is the surplus or deficit of that whole, in millions of euros
and as a percentage of GDP. It is the country's balance, not the central
government alone and not a regional government.
"""

from __future__ import annotations

import csv
import os
import re
import subprocess
import sys
import tempfile
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "data" / "spain.csv"
EU_OUT = ROOT / "data" / "spain-eu.csv"
DEBT_OUT = ROOT / "data" / "spain-debt.csv"
URL = (
    "https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/"
    "ContabilidadNacional/Publicaciones/Documents/AAPP_A/A_AAPP.xlsx"
)
CENTRAL_URL = (
    "https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/"
    "ContabilidadNacional/Publicaciones/Documents/AACC-A/A_AACC.xlsx"
)
EU_URL = (
    "https://commission.europa.eu/document/download/"
    "9f334bee-d097-4b68-8f93-2225eabeb631_en"
    "?filename=eu_budget_spending_and_revenue_2000-2025.xlsx"
)
DEBT_URL = "https://www.bde.es/webbe/es/estadisticas/compartido/datos/csv/be1112.csv"
USER_AGENT = "plan-mery/spain-downloader"
# Non-financial uses of all public administrations. Internal transfers cancel
# in the consolidated account, so they stay out. The sum is Eurostat's total
# expenditure for Spain.
SPENDING_CODES = (
    "D.1",
    "P.2",
    "D.29",
    "D.31",
    "D.39",
    "D.41",
    "D.42+..+D.45",
    "D.51",
    "D.62",
    "D.632",
    "D.71",
    "D.72",
    "D.74",
    "D.75",
    "D.76",
    "D.8",
    "P.51g",
    "P.52+P.53",
    "NP",
    "D.92 (exc. D.9_S.13)",
    "D.99 (exc. D.9_S.13)",
)
DEBT_FIELDS = {
    "DTNPDE2010_P00000_PS_APU": "total",
    "DTNPDE2010_F30000_PS_APU": "bonds",
    "DTNPDE2010_F31000_PS_APU": "bonds_short",
    "DTNPDE2010_F32000_PS_APU": "bonds_long",
    "DTNPDE2010_F40000_PS_APU": "loans",
}


def ensure_openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        pass
    venv = Path(tempfile.gettempdir()) / "plan-mery-xlsvenv"
    python = venv / "bin" / "python"
    if not python.exists():
        subprocess.check_call([sys.executable, "-m", "venv", str(venv)])
    if subprocess.run([str(python), "-c", "import openpyxl"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL).returncode != 0:
        subprocess.check_call([str(python), "-m", "pip", "install", "openpyxl"])
    if Path(sys.prefix).resolve() != venv.resolve():
        os.execv(str(python), [str(python), __file__, *sys.argv[1:]])
    import openpyxl
    return openpyxl


def year_of(cell) -> int | None:
    if isinstance(cell, bool) or cell in (None, ""):
        return None
    if isinstance(cell, (int, float)):
        return int(cell)
    match = re.match(r"(\d{4})", str(cell).strip())
    return int(match.group(1)) if match else None


def series(sheet) -> tuple[list[int], list[float | None]]:
    rows = list(sheet.iter_rows(values_only=True))
    header = None
    for row in rows:
        if row and row[0] == "CÓDIGO":
            header = row
            break
    if header is None:
        raise SystemExit("Spain accounts have no year header")
    years = []
    for cell in header[2:]:
        year = year_of(cell)
        if year is None:
            break
        years.append(year)
    values = None
    for row in rows:
        if row and row[0] == "B.9":
            values = [None if cell in (None, "") else float(cell) for cell in row[2:2 + len(years)]]
            break
    if values is None:
        raise SystemExit("Spain accounts have no B.9 balance")
    return years, values


def central_spending(sheet) -> dict[int, float]:
    """Total non-financial expenditure of the central administration, millions of euros."""
    rows = list(sheet.iter_rows(values_only=True))
    header = None
    for row in rows:
        if row and row[0] == "CÓDIGO":
            header = row
            break
    if header is None:
        raise SystemExit("Central accounts have no year header")
    years = []
    for cell in header[2:]:
        year = year_of(cell)
        if year is None:
            break
        years.append(year)
    for row in rows:
        name = row[1] if len(row) > 1 else None
        if isinstance(name, str) and name.strip().startswith("EMPLEOS NO FINANCIEROS"):
            found = {}
            for year, cell in zip(years, row[2:]):
                if cell in (None, ""):
                    continue
                found[year] = float(cell)
            if 2019 not in found or 2025 not in found:
                raise SystemExit("Central spending is missing 2019 or 2025")
            return found
    raise SystemExit("Central spending total is missing")


def spending(sheet, years: list[int]) -> list[float | None]:
    """Total expenditure, millions of euros, one value per year."""
    rows = list(sheet.iter_rows(values_only=True))
    section = ""
    found: dict[str, list] = {}
    for row in rows:
        name = row[1] if len(row) > 1 else None
        if isinstance(name, str) and name.strip() in ("EMPLEOS CORRIENTES", "EMPLEOS DE CAPITAL", "RECURSOS CORRIENTES"):
            section = name.strip()
        code = row[0]
        if section != "EMPLEOS CORRIENTES" and section != "EMPLEOS DE CAPITAL":
            continue
        if code == "B.9":
            break
        if code not in SPENDING_CODES:
            continue
        found[code] = [None if cell in (None, "") else float(cell) for cell in row[2:2 + len(years)]]
    missing = [code for code in SPENDING_CODES if code not in found]
    if missing:
        raise SystemExit(f"Spain spending is missing {missing[0]}")
    totals = []
    for index in range(len(years)):
        parts = [found[code][index] for code in SPENDING_CODES]
        totals.append(None if any(part is None for part in parts) else sum(parts))
    return totals


def spain_column(sheet) -> int:
    for row in sheet.iter_rows(values_only=True):
        for index, cell in enumerate(row):
            if isinstance(cell, str) and cell.strip() == "ES":
                return index
    raise SystemExit(f"EU budget sheet {sheet.title} has no Spain column")


def row_label(row) -> str:
    for cell in row:
        if isinstance(cell, str) and cell.strip():
            text = " ".join(cell.split())
            if not text[0].isdigit():
                return text
    return ""


def eu_year(sheet) -> tuple[float, float] | None:
    """EU spending in Spain, and Spain's payment into the EU budget, in millions."""
    column = spain_column(sheet)
    received = None
    recovery = 0.0
    own = None
    customs = 0.0
    sugar = 0.0
    includes_customs = False
    for row in sheet.iter_rows(values_only=True):
        label = row_label(row)
        value = row[column] if column < len(row) else None
        number = value if isinstance(value, (int, float)) and not isinstance(value, bool) else None
        if label == "TOTAL EXPENDITURE" and number is not None:
            received = float(number)
        elif label == "TOTAL NGEU" and number is not None:
            recovery = float(number)
        elif label in ("TOTAL own resources", "TOTAL Own resources") and number is not None:
            own = float(number)
        elif label == "Customs duties" and number is not None:
            customs = float(number)
        elif label == "Sugar levies" and number is not None:
            sugar = float(number)
        elif label.startswith("Traditional own resources"):
            includes_customs = True
    if received is None or own is None:
        return None
    paid = own if includes_customs else own + customs + sugar
    return received + recovery, paid


def eu_flows(book) -> list[tuple[int, float, float]]:
    flows = []
    for name in book.sheetnames:
        if not name.isdigit():
            continue
        parsed = eu_year(book[name])
        if parsed is None:
            continue
        received, paid = parsed
        flows.append((int(name), received, paid))
    if not flows:
        raise SystemExit("EU budget workbook has no Spain totals")
    return flows


def spain_debt(payload: bytes) -> list[tuple]:
    rows = list(csv.reader(payload.decode("latin-1").splitlines()))
    columns = {code: index for index, code in enumerate(rows[0]) if code in DEBT_FIELDS}
    if set(columns) != set(DEBT_FIELDS):
        raise SystemExit("Spain debt file is missing an instrument")
    found = []
    for row in rows[6:]:
        label = row[0].strip().strip('"')
        if not label.startswith("DIC "):
            continue
        year = int(label.split()[-1])
        values = {}
        for code, index in columns.items():
            raw = row[index].strip().strip('"')
            if raw in ("", "_"):
                values = {}
                break
            values[DEBT_FIELDS[code]] = float(raw)
        if values:
            found.append((year, values["total"], values["bonds"], values["bonds_short"], values["bonds_long"], values["loans"]))
    if not found:
        raise SystemExit("Spain debt file has no December stock")
    return found


def main() -> None:
    openpyxl = ensure_openpyxl()
    request = urllib.request.Request(URL, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=60) as response:
        payload = response.read()
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as handle:
        handle.write(payload)
        path = Path(handle.name)
    try:
        book = openpyxl.load_workbook(path, data_only=True, read_only=True)
        euros_years, euros = series(book["Tabla1a"])
        percent_years, percents = series(book["Tabla1b"])
        spent = spending(book["Tabla1a"], euros_years)
    finally:
        path.unlink(missing_ok=True)
    if euros_years != percent_years:
        raise SystemExit("Spain euro years and percent years differ")
    request = urllib.request.Request(CENTRAL_URL, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=60) as response:
        central_payload = response.read()
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as handle:
        handle.write(central_payload)
        central_path = Path(handle.name)
    try:
        central_book = openpyxl.load_workbook(central_path, data_only=True, read_only=True)
        central = central_spending(central_book["Tabla1a"])
    finally:
        central_path.unlink(missing_ok=True)
    request = urllib.request.Request(EU_URL, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=90) as response:
        eu_payload = response.read()
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as handle:
        handle.write(eu_payload)
        eu_path = Path(handle.name)
    try:
        eu_book = openpyxl.load_workbook(eu_path, data_only=True, read_only=True)
        flows = eu_flows(eu_book)
    finally:
        eu_path.unlink(missing_ok=True)
    request = urllib.request.Request(DEBT_URL, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=60) as response:
        debt_rows = spain_debt(response.read())
    OUT.parent.mkdir(exist_ok=True)
    with OUT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(["year", "million", "percent", "spent", "central"])
        written = 0
        for year, million, percent, expenditure in zip(euros_years, euros, percents, spent):
            state = central.get(year)
            if million is None or percent is None or expenditure is None or state is None:
                continue
            writer.writerow((year, million, percent, round(expenditure, 3), round(state, 3)))
            written += 1
    with EU_OUT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(["year", "received", "paid"])
        for year, received, paid in flows:
            writer.writerow((year, round(received, 3), round(paid, 3)))
    with DEBT_OUT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(["year", "total", "bonds", "bonds_short", "bonds_long", "loans"])
        writer.writerows(debt_rows)
    print(f"Wrote {OUT} ({written} years), {EU_OUT} ({len(flows)} years), {DEBT_OUT} ({len(debt_rows)} years)")


if __name__ == "__main__":
    main()
