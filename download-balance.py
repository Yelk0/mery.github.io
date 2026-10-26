#!/usr/bin/env python3
"""Download each community's national-accounts balance and write data/balance.csv.

The Intervención General publishes the annual non-financial accounts of the
regional administration, one workbook per year. The B.9 line is the
capacidad (+) or necesidad (−) de financiación, in millions of euros and as
a percentage of that community's GDP. Ceuta and Melilla are not in the file.
"""

from __future__ import annotations

import csv
import os
import subprocess
import sys
import tempfile
import unicodedata
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "data" / "balance.csv"
PAGE = (
    "https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/"
    "ContabilidadNacional/Publicaciones/Paginas/ianofinancierasCA.aspx"
)
FILE = (
    "https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/"
    "ContabilidadNacional/Publicaciones/Documents/CCAA-A/A_CCAA_Det_{year}.xlsx"
)
USER_AGENT = "plan-mery/balance-downloader"
FIRST_YEAR = 2000
LAST_YEAR = 2025

NAMES = {
    "andalucia": "andalucia",
    "aragon": "aragon",
    "asturias": "asturias",
    "baleares": "baleares",
    "canarias": "canarias",
    "cantabria": "cantabria",
    "castilla-la mancha": "castilla-la-mancha",
    "castilla y leon": "castilla-y-leon",
    "cataluna": "cataluna",
    "extremadura": "extremadura",
    "galicia": "galicia",
    "comunidad de madrid": "madrid",
    "region de murcia": "murcia",
    "comunidad foral de navarra": "navarra",
    "la rioja": "la-rioja",
    "comunitat valenciana": "valencia",
    "comunitat valencia": "valencia",
    "pais vasco": "pais-vasco",
}


def ensure_openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        pass
    venv = Path(tempfile.gettempdir()) / "plan-mery-xlsvenv"
    python = venv / "bin" / "python"
    if not python.exists():
        subprocess.check_call([sys.executable, "-m", "venv", str(venv)])
    if subprocess.run([str(python), "-c", "import openpyxl"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL).returncode != 0:
        subprocess.check_call([str(python), "-m", "pip", "install", "openpyxl"])
    if Path(sys.prefix).resolve() != venv.resolve():
        os.execv(str(python), [str(python), __file__, *sys.argv[1:]])
    import openpyxl
    return openpyxl


def fold(value: str) -> str:
    text = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return " ".join(text.lower().split())


def download(year: int) -> bytes:
    request = urllib.request.Request(FILE.format(year=year), headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=120) as response:
        payload = response.read()
    if payload[:2] != b"PK":
        raise SystemExit(f"{year} did not return a workbook")
    return payload


def read_year(openpyxl, year: int, payload: bytes) -> list[tuple]:
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as handle:
        handle.write(payload)
        path = Path(handle.name)
    try:
        book = openpyxl.load_workbook(path, data_only=True, read_only=True)
        millions = line(book["Tabla1a"])
        percent = line(book["Tabla1b"])
    finally:
        path.unlink(missing_ok=True)
    rows = []
    for name, million in millions.items():
        slug = NAMES.get(name)
        if slug is None or name not in percent:
            continue
        rows.append((slug, year, million, percent[name]))
    if len(rows) < 17:
        raise SystemExit(f"{year} has {len(rows)} communities, expected 17")
    return rows


def line(sheet) -> dict[str, float]:
    header_row = None
    for row in range(1, 12):
        concept = sheet.cell(row, 2).value
        if concept and "CONCEPTO" in str(concept).upper():
            header_row = row
            break
    if header_row is None:
        raise SystemExit("Workbook has no concept header")
    columns = {}
    for column in range(3, sheet.max_column + 1):
        name = sheet.cell(header_row, column).value
        if name:
            columns[column] = fold(str(name))
    balance_row = None
    for row in range(header_row + 1, sheet.max_row + 1):
        code = sheet.cell(row, 1).value
        concept = str(sheet.cell(row, 2).value or "")
        if code == "B.9" or concept.startswith("CAPACIDAD"):
            balance_row = row
            break
    if balance_row is None:
        raise SystemExit("Workbook has no B.9 balance")
    values = {}
    for column, name in columns.items():
        if name.startswith("total"):
            continue
        cell = sheet.cell(balance_row, column).value
        if cell is None or cell == "":
            continue
        values[name] = float(cell)
    return values


def main() -> None:
    openpyxl = ensure_openpyxl()
    OUT.parent.mkdir(exist_ok=True)
    rows = []
    for year in range(FIRST_YEAR, LAST_YEAR + 1):
        rows.extend(read_year(openpyxl, year, download(year)))
        print(f"{year}: {sum(1 for row in rows if row[1] == year)} communities")
    rows.sort(key=lambda row: (row[0], row[1]))
    with OUT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(["slug", "year", "million", "percent"])
        writer.writerows(rows)
    print(f"Wrote {OUT} ({len(rows)} balances) from {PAGE}")


if __name__ == "__main__":
    main()
