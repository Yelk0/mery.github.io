#!/usr/bin/env python3
"""Download settled transfers between each community and the State.

The Ministry of Finance publishes the consolidated liquidation. On the revenue
side, chapter lines that start with 40 or 70 are money moving between the
State and the community. A positive line is money the State transfers in. A
negative line is money the community pays into the financing system. The
share of State taxes that the community only passes on to town councils stays
out. On the spending side, chapters 40 and 70 paid to the State are added to
what the community pays back. The files start in 2002 and currently end in
2024. Ceuta and Melilla are requested and left blank when the workbook is
missing.
"""

from __future__ import annotations

import csv
import http.cookiejar
import os
import re
import subprocess
import sys
import tempfile
import time
import unicodedata
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "data" / "transfers.csv"
PAGE = "https://serviciostelematicosext.hacienda.gob.es/SGCIEF/PublicacionLiquidaciones/aspx/menuInicio.aspx"
REVENUE = (
    "https://serviciostelematicosext.hacienda.gob.es/SGCIEF/PublicacionLiquidaciones/"
    "aspx/DescargaIngresosDC.aspx?cdcdad={code}&ano={year}"
)
SPENDING = (
    "https://serviciostelematicosext.hacienda.gob.es/SGCIEF/PublicacionLiquidaciones/"
    "aspx/DescargaGastosDC.aspx?cdcdad={code}&ano={year}"
)
USER_AGENT = "plan-mery/transfer-downloader"
FIRST_YEAR = 2002
LAST_YEAR = 2024
CODE = re.compile(r"(\d{2,4})")
COMMUNITIES = {
    "01": "andalucia",
    "02": "aragon",
    "03": "asturias",
    "04": "baleares",
    "05": "canarias",
    "06": "cantabria",
    "07": "castilla-y-leon",
    "08": "castilla-la-mancha",
    "09": "cataluna",
    "10": "extremadura",
    "11": "galicia",
    "12": "madrid",
    "13": "murcia",
    "14": "navarra",
    "15": "pais-vasco",
    "16": "la-rioja",
    "17": "valencia",
    "18": "ceuta",
    "19": "melilla",
}


def ensure_openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        pass
    venv = Path(tempfile.gettempdir()) / "plan-mery-xlsvenv"
    python = venv / "bin" / "python"
    if not python.exists():
        subprocess.check_call([sys.executable, "-m", "venv", str(venv)])
    if subprocess.run([str(python), "-c", "import openpyxl"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL).returncode != 0:
        subprocess.check_call([str(python), "-m", "pip", "install", "openpyxl"])
    if Path(sys.prefix).resolve() != venv.resolve():
        os.execv(str(python), [str(python), __file__, *sys.argv[1:]])
    import openpyxl
    return openpyxl


def fold(value: str) -> str:
    text = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return " ".join(text.lower().split())


def opener() -> urllib.request.OpenerDirector:
    jar = http.cookiejar.CookieJar()
    agent = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
    request = urllib.request.Request(PAGE, headers={"User-Agent": USER_AGENT})
    with agent.open(request, timeout=60) as response:
        response.read()
    return agent


def fetch(agent: urllib.request.OpenerDirector, url: str) -> bytes | None:
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    try:
        with agent.open(request, timeout=90) as response:
            payload = response.read()
    except urllib.error.HTTPError:
        return None
    if payload[:2] != b"PK":
        return None
    return payload


def sheet_rows(openpyxl, payload: bytes) -> list[tuple]:
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as handle:
        handle.write(payload)
        path = Path(handle.name)
    try:
        book = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheet = book[book.sheetnames[0]]
        return list(sheet.iter_rows(values_only=True))
    finally:
        path.unlink(missing_ok=True)


def coded_rows(rows: list[tuple]) -> list[tuple[str, str, float | None]]:
    found = []
    for row in rows:
        if not row or row[0] is None:
            continue
        label = str(row[0]).strip()
        match = CODE.search(label)
        if match is None:
            continue
        value = row[1] if len(row) > 1 else None
        amount = None if value in (None, "") else float(value)
        found.append((match.group(1), label, amount))
    return found


def leaves(rows: list[tuple[str, str, float | None]]) -> list[tuple[str, str, float | None]]:
    chosen = []
    for index, item in enumerate(rows):
        code = item[0]
        child = False
        for other, _, _ in rows[index + 1:]:
            if other.startswith(code) and len(other) > len(code):
                child = True
                break
            if not other.startswith(code):
                break
        if not child:
            chosen.append(item)
    return chosen


def state_flows(revenue_rows: list[tuple], spending_rows: list[tuple]) -> tuple[float, float] | None:
    """Miles de euros. None when the workbook has no figures on these lines."""
    received = 0.0
    paid = 0.0
    seen = False
    for code, _, amount in leaves(coded_rows(revenue_rows)):
        if amount is None or code == "4002":
            continue
        if not (code.startswith("40") or code.startswith("70")):
            continue
        if len(code) < 3:
            continue
        seen = True
        if amount >= 0:
            received += amount
        else:
            paid += -amount
    for code, label, amount in coded_rows(spending_rows):
        if amount is None or code not in ("40", "70"):
            continue
        if "estado" not in fold(label):
            continue
        seen = True
        paid += amount
    if not seen:
        return None
    return received, paid


def saved_keys() -> set[tuple[str, int]]:
    if not OUT.exists():
        return set()
    with OUT.open(encoding="utf-8", newline="") as handle:
        return {(row["slug"], int(row["year"])) for row in csv.DictReader(handle, delimiter=";")}


def main() -> None:
    openpyxl = ensure_openpyxl()
    agent = opener()
    OUT.parent.mkdir(exist_ok=True)
    done = saved_keys()
    fresh = not OUT.exists()
    handle = OUT.open("a", encoding="utf-8", newline="")
    writer = csv.writer(handle, delimiter=";")
    if fresh:
        writer.writerow(["slug", "year", "received", "paid"])
        handle.flush()
    added = 0
    try:
        for code, slug in COMMUNITIES.items():
            for year in range(FIRST_YEAR, LAST_YEAR + 1):
                if (slug, year) in done:
                    continue
                revenue = fetch(agent, REVENUE.format(code=code, year=year))
                time.sleep(0.15)
                spending = fetch(agent, SPENDING.format(code=code, year=year))
                time.sleep(0.15)
                if revenue is None or spending is None:
                    print(f"{slug} {year}: missing", flush=True)
                    continue
                flows = state_flows(sheet_rows(openpyxl, revenue), sheet_rows(openpyxl, spending))
                if flows is None:
                    print(f"{slug} {year}: missing", flush=True)
                    continue
                received, paid = flows
                writer.writerow((slug, year, round(received, 2), round(paid, 2)))
                handle.flush()
                added += 1
                print(f"{slug} {year}: received {received / 1000:.0f} paid {paid / 1000:.0f}", flush=True)
    finally:
        handle.close()
    print(f"Wrote {OUT} ({len(done) + added} years, {added} new)")


if __name__ == "__main__":
    main()
