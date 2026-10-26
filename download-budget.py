#!/usr/bin/env python3
"""Download settled regional budgets by spending policy and write data/budget.csv.

The Ministry of Finance publishes the consolidated liquidation of each
community, classified by function and by chapter. Health, education, care,
housing, and transport come from the policy rows. Interest is the financial
chapter of spending. Repayment of principal and the pensions policy stay out
of the total. Every other spending policy is kept on its own. The liquidation
database begins in 2002 and currently ends in 2024. Before the policy
classification, social protection and social promotion are the care line,
because pensions were not published apart.
"""

from __future__ import annotations

import csv
import http.cookiejar
import os
import re
import subprocess
import sys
import tempfile
import time
import unicodedata
import urllib.request
from pathlib import Path

from mery.selection import limit

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "data" / "budget.csv"
PAGE = "https://serviciostelematicosext.hacienda.gob.es/SGCIEF/PublicacionLiquidaciones/aspx/menuInicio.aspx"
FILE = (
    "https://serviciostelematicosext.hacienda.gob.es/SGCIEF/PublicacionLiquidaciones/"
    "aspx/DescargaFuncionalCapDC.aspx?cdcdad={code}&ano={year}"
)
USER_AGENT = "plan-mery/budget-downloader"
FIRST_YEAR = 2002
LAST_YEAR = 2024
DETAIL = re.compile(r"^(\d{2})\s*\.")

COMMUNITIES = {
    "01": "andalucia",
    "02": "aragon",
    "03": "asturias",
    "04": "baleares",
    "05": "canarias",
    "06": "cantabria",
    "07": "castilla-y-leon",
    "08": "castilla-la-mancha",
    "09": "cataluna",
    "10": "extremadura",
    "11": "galicia",
    "12": "madrid",
    "13": "murcia",
    "14": "navarra",
    "15": "pais-vasco",
    "16": "la-rioja",
    "17": "valencia",
    "18": "ceuta",
    "19": "melilla",
}
TARGETS = (
    "justice", "defence", "security", "foreign", "employment", "unemployment", "benefits",
    "social_security", "culture", "agriculture", "farm_works", "industry", "energy",
    "commerce", "tourism", "research", "economic", "leadership", "general", "statistics",
    "tax", "transfers", "communications", "community", "other_social",
)
NEW = {
    "11": "justice",
    "12": "defence",
    "13": "security",
    "14": "foreign",
    "21": "pensions",
    "22": "benefits",
    "23": "care",
    "24": "employment",
    "25": "unemployment",
    "26": "housing",
    "29": "social_security",
    "31": "health",
    "32": "education",
    "33": "culture",
    "41": "agriculture",
    "42": "industry",
    "43": "commerce",
    "44": "transport",
    "45": "transport",
    "46": "research",
    "49": "economic",
    "91": "leadership",
    "92": "general",
    "93": "tax",
    "94": "transfers",
    "95": "debt",
}
OLD = {
    "01": "debt",
    "11": "leadership",
    "12": "general",
    "13": "foreign",
    "14": "justice",
    "22": "security",
    "31": "care",
    "32": "care",
    "41": "health",
    "42": "education",
    "43": "housing",
    "44": "community",
    "45": "culture",
    "46": "other_social",
    "51": "transport",
    "52": "communications",
    "53": "farm_works",
    "54": "research",
    "55": "statistics",
    "61": "economic",
    "62": "commerce",
    "63": "tax",
    "71": "agriculture",
    "72": "industry",
    "73": "energy",
    "74": "energy",
    "75": "tourism",
    "91": "transfers",
}
COLUMNS = ("health", "education", "care", "housing", "transport", "interest", *TARGETS, "rest", "spent")


def ensure_openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        pass
    venv = Path(tempfile.gettempdir()) / "plan-mery-xlsvenv"
    python = venv / "bin" / "python"
    if not python.exists():
        subprocess.check_call([sys.executable, "-m", "venv", str(venv)])
    if subprocess.run([str(python), "-c", "import openpyxl"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL).returncode != 0:
        subprocess.check_call([str(python), "-m", "pip", "install", "openpyxl"])
    if Path(sys.prefix).resolve() != venv.resolve():
        os.execv(str(python), [str(python), __file__, *sys.argv[1:]])
    import openpyxl
    return openpyxl


def fold(value: str) -> str:
    text = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return " ".join(text.lower().split())


def opener() -> urllib.request.OpenerDirector:
    jar = http.cookiejar.CookieJar()
    agent = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
    request = urllib.request.Request(PAGE, headers={"User-Agent": USER_AGENT})
    with agent.open(request, timeout=60) as response:
        response.read()
    return agent


def fetch(agent: urllib.request.OpenerDirector, code: str, year: int) -> bytes:
    target = Path(tempfile.gettempdir()) / f"plan-mery-ccaa-{code}-{year}.xlsx"
    if target.exists() and target.stat().st_size > 1000:
        return target.read_bytes()
    request = urllib.request.Request(FILE.format(code=code, year=year), headers={"User-Agent": USER_AGENT})
    with agent.open(request, timeout=90) as response:
        payload = response.read()
    if payload[:2] != b"PK":
        raise RuntimeError(f"{code} {year} did not return a workbook")
    target.write_bytes(payload)
    return payload


def amount(value) -> float:
    if value is None or value == "":
        return 0.0
    return float(value)


def chapter_key(label: str) -> str | None:
    if label.startswith("gastos de personal"):
        return "staff"
    if label.startswith("gastos corrientes"):
        return "goods"
    if "transferenc" in label and "capital" in label:
        return "capital_transfers"
    if "transferenc" in label and "corrient" in label:
        return "current_transfers"
    if label.startswith("inversiones reales"):
        return "investment"
    if label.startswith("activos financieros"):
        return "assets"
    if "contingencia" in label:
        return "contingency"
    return None


def general_chapters(header: list[str], row: tuple, target: float) -> dict[str, float]:
    parts: dict[str, float] = {}
    for index, label in enumerate(header):
        key = chapter_key(label)
        if key is None or index >= len(row):
            continue
        parts[key] = parts.get(key, 0.0) + amount(row[index])
    gap = target - sum(parts.values())
    if abs(gap) > 1:
        parts["general_other"] = parts.get("general_other", 0.0) + gap
    return {key: round(value, 2) for key, value in parts.items() if value > 1}


def read_book(openpyxl, payload: bytes) -> tuple[dict[str, float], dict[str, float]]:
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as handle:
        handle.write(payload)
        path = Path(handle.name)
    try:
        book = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheet = book[book.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        path.unlink(missing_ok=True)
    header_index = None
    for index, row in enumerate(rows):
        labels = [fold(str(cell)) for cell in row if cell]
        if any(label.startswith("gastos financieros") for label in labels):
            header_index = index
            break
    if header_index is None:
        raise RuntimeError("Workbook has no financial-chapter column")
    header = [fold(str(cell)) if cell else "" for cell in rows[header_index]]
    financial = next(index for index, label in enumerate(header) if label.startswith("gastos financieros"))
    liabilities = next(index for index, label in enumerate(header) if label.startswith("pasivos financieros"))
    total = next(index for index, label in enumerate(header) if label.startswith("total"))
    details = []
    for row in rows[header_index + 1:]:
        if not row or row[0] is None:
            continue
        label = str(row[0]).strip()
        match = DETAIL.match(label)
        if not match:
            continue
        details.append((match.group(1), fold(label.split(".", 1)[1]), row))
    if not details:
        raise RuntimeError("Workbook has no policy rows")
    labels = {code: name for code, name, _row in details}
    if "sanidad" in labels.get("31", ""):
        buckets = NEW
    elif "sanidad" in labels.get("41", ""):
        buckets = OLD
    else:
        raise RuntimeError("Workbook uses an unknown policy classification")
    spent = {name: 0.0 for name in COLUMNS}
    unknown = []
    chapters: dict[str, float] = {}
    for code, name, row in details:
        bucket = buckets.get(code)
        if bucket is None:
            unknown.append(f"{code} {name}")
            continue
        whole = amount(row[total] if total < len(row) else None)
        interest = amount(row[financial] if financial < len(row) else None)
        principal = amount(row[liabilities] if liabilities < len(row) else None)
        if bucket == "pensions":
            continue
        if bucket != "debt":
            net = whole - interest - principal
            spent[bucket] += net
            if bucket == "general":
                chapters = general_chapters(header, row, net)
        spent["interest"] += interest
    if unknown:
        raise RuntimeError(f"Unknown policy codes: {unknown}")
    spent["spent"] = sum(spent[name] for name in COLUMNS if name != "spent")
    return {name: round(spent[name], 2) for name in COLUMNS}, chapters


def main() -> None:
    openpyxl = ensure_openpyxl()
    OUT.parent.mkdir(exist_ok=True)
    agent = opener()
    rows = []
    part_rows = {}
    for year in range(FIRST_YEAR, LAST_YEAR + 1):
        for code, slug in limit(COMMUNITIES).items():
            payload = None
            for attempt in range(3):
                try:
                    payload = fetch(agent, code, year)
                    break
                except Exception as error:
                    print(f"retry {slug} {year}: {error}", file=sys.stderr)
                    time.sleep(1.5)
                    agent = opener()
            if payload is None:
                raise SystemExit(f"Could not download {slug} {year}")
            try:
                values, chapters = read_book(openpyxl, payload)
            except Exception as error:
                raise SystemExit(f"Could not read {slug} {year}: {error}") from error
            if values["spent"] <= 0:
                print(f"skip {slug} {year}: the liquidation has no spending", file=sys.stderr)
                time.sleep(0.05)
                continue
            rows.append((slug, year, *(values[name] for name in COLUMNS)))
            part_rows[(slug, year)] = chapters
            time.sleep(0.05)
        print(f"{year}: {sum(1 for row in rows if row[1] == year)} communities")
    rows.sort(key=lambda row: (row[0], row[1]))
    with OUT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(["slug", "year", *COLUMNS])
        writer.writerows(rows)
    import importlib.util
    spec = importlib.util.spec_from_file_location("state_budget", ROOT / "download-state-budget.py")
    state_budget = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(state_budget)
    state_budget.save_parts(part_rows)
    print(f"Wrote {OUT} ({len(rows)} budgets) from {PAGE}")


if __name__ == "__main__":
    main()
